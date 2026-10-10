"""Ingesta de la base sintética del reto (Lulo Bank · Reto 19) para el motor v3.

Lee el Excel original, valida el esquema y la calidad de cada pestaña y escribe
CSV normalizados en lulo_app/dataset/, junto con un manifiesto auditable y un
reporte de calidad legible.

Reglas de la ingesta:
  * Nada se corrige en silencio. Cada anomalía queda en el manifiesto con su
    conteo, y lo que se excluye se escribe aparte con la razón.
  * Los errores de esquema detienen la ingesta (código de salida 1).
  * La salida es determinista: el mismo Excel produce byte a byte los mismos
    archivos, así que las pruebas pueden verificar que lo publicado coincide.

Uso:
    python scripts/ingest_reto.py                      # usa data/raw/Base_Datos_Reto_Universidad.xlsx
    python scripts/ingest_reto.py --src otro.xlsx --out /tmp/salida
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - mensaje claro en vez de un traceback
    sys.exit("Falta openpyxl: pip install -r requirements-dev.txt")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SRC = ROOT / "data" / "raw" / "Base_Datos_Reto_Universidad.xlsx"
DEFAULT_OUT = ROOT / "lulo_app" / "dataset"
INGEST_VERSION = "1.0.0"

HEADER_ROW = 3  # fila 1: título de la pestaña, fila 2: vacía, fila 3: encabezados

SCHEMAS: dict[str, list[str]] = {
    "Clientes": ["cliente_id", "rango_edad", "ocupacion", "antiguedad_meses", "segmento", "saldo_promedio_cop",
                 "ingreso_mensual_cop", "credito_activo", "cajitas_activas", "cdt_activo", "seguros_vigentes",
                 "autoriza_ofertas"],
    "Movimientos": ["cliente_id", "fecha", "monto_cop", "canal", "categoria", "mcc", "tipo_movimiento"],
    "Pagos_Seguros": ["cliente_id", "fecha_pago", "ramo", "monto_cop", "fecha_vencimiento", "proveedor"],
    "Cajitas": ["cliente_id", "fecha_abono", "tipo_cajita", "monto_abono_cop"],
    "Uso_App": ["cliente_id", "sesiones_30d", "visitas_seguros_30d", "notificaciones_enviadas_30d", "abiertas_30d",
                "ignoradas_30d", "max_contactos_mes", "latencia_eventos"],
    "Seguros_Mensual": ["mes", "ramo", "polizas_emitidas", "renovaciones", "prima_promedio_cop", "comision_pct",
                        "persistencia_6m", "persistencia_12m", "cancelacion_pct"],
    "Elegibilidad": ["ramo", "edad", "ocupacion", "requisito_simulado", "periodicidad"],
    "Campanas_Previas": ["campana", "canal", "elegibles", "impactados", "interesados", "ventas",
                         "conversion_impactados"],
    "Juridico_Entrega": ["tema", "supuesto_simulado", "nota"],
}

ENUMS = {
    "rango_edad": {"18-25", "26-35", "36-45", "46-55", "56+"},
    "ocupacion": {"Empleado", "Independiente", "Estudiante", "Pensionado"},
    "segmento": {"Pro", "No Pro"},
    "si_no": {"Sí", "No"},
    "canal": {"Transferencia", "Tarjeta débito", "PSE", "QR"},
    "tipo_movimiento": {"Crédito", "Débito"},
    "ramo": {"SOAT", "Vida", "Desempleo", "Hogar", "Mascotas", "Viajes"},
    "tipo_cajita": {"Emergencias", "Viajes", "Otras", "Vivienda", "Mascotas"},
    "latencia_eventos": {"<1 min", "1-5 min", "5-30 min", "Mismo día"},
}
CLIENT_ID = re.compile(r"^C\d{5}$")


# --------------------------------------------------------------------------- utilidades
class SchemaError(Exception):
    pass


@dataclass
class Report:
    checks: list[dict] = field(default_factory=list)

    def add(self, sheet: str, check: str, status: str, detail: str, count: int | None = None) -> None:
        assert status in {"ok", "aviso", "error"}
        self.checks.append({"pestana": sheet, "chequeo": check, "estado": status, "detalle": detail,
                            **({"conteo": count} if count is not None else {})})

    @property
    def errors(self) -> list[dict]:
        return [c for c in self.checks if c["estado"] == "error"]


def _iso(v) -> str:
    if isinstance(v, datetime):
        if (v.hour, v.minute, v.second, v.microsecond) != (0, 0, 0, 0):
            raise SchemaError(f"Fecha con hora inesperada: {v!r}")
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v.strip()):
        return v.strip()
    raise SchemaError(f"No es una fecha: {v!r}")


def _int(v) -> int:
    if isinstance(v, bool):
        raise SchemaError(f"No es un entero: {v!r}")
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and re.fullmatch(r"-?\d+", v.strip()):
        return int(v.strip())
    raise SchemaError(f"No es un entero: {v!r}")


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    raise SchemaError(f"No es un número: {v!r}")


def _fmt_num(v: float | None) -> str:
    if v is None:
        return ""
    return repr(round(v, 6)).removesuffix(".0") if float(v).is_integer() else repr(round(v, 6))


def _str(v) -> str:
    if v is None:
        raise SchemaError("Valor vacío")
    return str(v).strip()


def read_sheet(wb, name: str, expected: list[str] | None = None) -> list[tuple[int, dict]]:
    """Devuelve [(fila_excel, {columna: valor})] validando los encabezados exactos."""
    expected = expected or SCHEMAS[name]
    if name not in wb.sheetnames:
        raise SchemaError(f"Falta la pestaña «{name}»")
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < HEADER_ROW:
        raise SchemaError(f"«{name}» no tiene encabezados en la fila {HEADER_ROW}")
    header = [c for c in rows[HEADER_ROW - 1] if c is not None]
    if header != expected:
        raise SchemaError(f"Encabezados de «{name}» no coinciden.\n  esperado: {expected}\n  recibido: {header}")
    out = []
    for i, row in enumerate(rows[HEADER_ROW:], start=HEADER_ROW + 1):
        vals = list(row[: len(header)])
        if all(v is None for v in vals):
            continue
        out.append((i, dict(zip(header, vals))))
    return out


def write_csv(path: Path, header: list[str], rows: list[list]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    w.writerows(rows)
    data = buf.getvalue().encode("utf-8")
    path.write_bytes(data)
    return hashlib.sha256(data).hexdigest()


def key_values(wb, name: str) -> dict[str, str]:
    out = {}
    for row in wb[name].iter_rows(values_only=True):
        if row and row[0] is not None and len(row) > 1 and row[1] is not None:
            out[str(row[0]).strip()] = row[1]
    return out


# --------------------------------------------------------------------------- ingesta
def ingest(src: Path, out: Path) -> dict:
    rep = Report()
    raw = src.read_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)

    # --- fecha de corte: dos fuentes en el archivo que deben coincidir
    jur = {r["tema"]: r for _, r in read_sheet(wb, "Juridico_Entrega")}
    corte = _iso(str(jur["Fecha de corte"]["supuesto_simulado"]).strip())
    periodo = str(key_values(wb, "LEEME").get("Periodo", ""))
    m = re.fullmatch(r"(\d{4}-\d{2}-\d{2}) a (\d{4}-\d{2}-\d{2})", periodo.strip())
    if not m:
        raise SchemaError(f"LEEME «Periodo» no tiene el formato esperado: {periodo!r}")
    inicio, fin = m.groups()
    if fin != corte:
        raise SchemaError(f"La fecha de corte de Juridico_Entrega ({corte}) no coincide con LEEME ({fin})")
    rep.add("LEEME/Juridico_Entrega", "fecha de corte", "ok", f"Periodo {inicio} a {corte}; ambas pestañas coinciden")
    leeme = key_values(wb, "LEEME")

    # --- Clientes
    clientes_rows = read_sheet(wb, "Clientes")
    clientes, ids = [], set()
    for fila, r in clientes_rows:
        cid = _str(r["cliente_id"])
        if not CLIENT_ID.match(cid):
            raise SchemaError(f"Clientes fila {fila}: cliente_id inválido {cid!r}")
        if cid in ids:
            raise SchemaError(f"Clientes fila {fila}: cliente_id duplicado {cid}")
        ids.add(cid)
        for col, enum in [("rango_edad", "rango_edad"), ("ocupacion", "ocupacion"), ("segmento", "segmento"),
                          ("credito_activo", "si_no"), ("cajitas_activas", "si_no"), ("cdt_activo", "si_no"),
                          ("autoriza_ofertas", "si_no")]:
            if _str(r[col]) not in ENUMS[enum]:
                raise SchemaError(f"Clientes fila {fila}: {col}={r[col]!r} fuera de {sorted(ENUMS[enum])}")
        seg = _str(r["seguros_vigentes"])
        tokens = [] if seg == "Ninguno" else [t.strip() for t in seg.split(",")]
        bad = [t for t in tokens if t not in ENUMS["ramo"]]
        if bad:
            raise SchemaError(f"Clientes fila {fila}: ramo desconocido en seguros_vigentes: {bad}")
        clientes.append([cid, _str(r["rango_edad"]), _str(r["ocupacion"]), _int(r["antiguedad_meses"]),
                         _str(r["segmento"]), _int(r["saldo_promedio_cop"]), _int(r["ingreso_mensual_cop"]),
                         _str(r["credito_activo"]), _str(r["cajitas_activas"]), _str(r["cdt_activo"]),
                         "|".join(tokens), _str(r["autoriza_ofertas"])])
    declared = _int(leeme.get("Clientes sintéticos", -1))
    rep.add("Clientes", "conteo vs LEEME", "ok" if declared == len(clientes) else "error",
            f"{len(clientes)} filas; LEEME declara {declared}", len(clientes))

    def fk(sheet: str, fila: int, cid: str) -> str:
        if cid not in ids:
            raise SchemaError(f"{sheet} fila {fila}: cliente_id {cid!r} no existe en Clientes")
        return cid

    excluded: list[list] = []

    # --- Movimientos
    movs, mcc_by_cat, n_after = [], defaultdict(set), Counter()
    credit_cats = Counter()
    for fila, r in read_sheet(wb, "Movimientos"):
        cid = fk("Movimientos", fila, _str(r["cliente_id"]))
        f = _iso(r["fecha"])
        monto = _int(r["monto_cop"])
        if monto <= 0:
            raise SchemaError(f"Movimientos fila {fila}: monto no positivo {monto}")
        canal, cat, tipo = _str(r["canal"]), _str(r["categoria"]), _str(r["tipo_movimiento"])
        mcc = _str(r["mcc"])
        if canal not in ENUMS["canal"] or tipo not in ENUMS["tipo_movimiento"]:
            raise SchemaError(f"Movimientos fila {fila}: canal/tipo inválido {canal!r}/{tipo!r}")
        if f < inicio:
            raise SchemaError(f"Movimientos fila {fila}: fecha {f} antes del inicio del periodo {inicio}")
        mcc_by_cat[cat].add(mcc)
        if tipo == "Crédito":
            credit_cats[cat] += 1
        row = [f"M{fila:06d}", cid, f, monto, canal, cat, mcc, tipo]
        if f > corte:
            n_after[cat] += 1
            excluded.append(["Movimientos", fila, cid, f, f"fecha posterior al corte {corte}", cat])
            continue
        movs.append(row)
    multi = {c: sorted(v) for c, v in mcc_by_cat.items() if len(v) > 1}
    rep.add("Movimientos", "categoría ↔ MCC 1:1", "ok" if not multi else "error",
            "cada categoría tiene un solo MCC" if not multi else f"categorías con varios MCC: {multi}")
    rep.add("Movimientos", "créditos (ingresos) por categoría", "ok",
            "; ".join(f"{k}: {v}" for k, v in sorted(credit_cats.items())), sum(credit_cats.values()))
    total_after = sum(n_after.values())
    rep.add("Movimientos", "fechas posteriores al corte", "aviso" if total_after else "ok",
            (f"{total_after} movimientos después de {corte} excluidos ("
             + ", ".join(f"{k}: {v}" for k, v in n_after.most_common()) + ")") if total_after else "ninguno",
            total_after)
    rep.add("Movimientos", "combustible", "aviso" if not any("combust" in c.lower() for c in mcc_by_cat) else "ok",
            "no existe categoría de combustible (MCC 5541/5542); la señal de vehículo solo cuenta con peajes y SOAT")

    # --- Pagos_Seguros
    pagos = []
    for fila, r in read_sheet(wb, "Pagos_Seguros"):
        cid = fk("Pagos_Seguros", fila, _str(r["cliente_id"]))
        fp, fv = _iso(r["fecha_pago"]), _iso(r["fecha_vencimiento"])
        ramo = _str(r["ramo"])
        if ramo not in ENUMS["ramo"]:
            raise SchemaError(f"Pagos_Seguros fila {fila}: ramo {ramo!r}")
        if fv <= fp:
            raise SchemaError(f"Pagos_Seguros fila {fila}: vence antes de pagarse")
        if fp > corte:
            excluded.append(["Pagos_Seguros", fila, cid, fp, f"fecha posterior al corte {corte}", ramo])
            continue
        pagos.append([f"P{fila:05d}", cid, fp, ramo, _int(r["monto_cop"]), fv, _str(r["proveedor"])])
    vig = {c[0]: set(c[10].split("|")) - {""} for c in clientes}
    overlap = Counter(); per_ramo = Counter()
    for p in pagos:
        per_ramo[p[3]] += 1
        overlap[p[3]] += p[3] in vig[p[1]]
    rep.add("Pagos_Seguros", "coincidencia con seguros_vigentes", "aviso",
            "porcentaje de pagos cuyo ramo aparece en seguros_vigentes del cliente: "
            + ", ".join(f"{k} {overlap[k]}/{n} ({overlap[k] / n:.0%})" for k, n in sorted(per_ramo.items()))
            + ". Lectura propuesta: vigentes = con Lulo; pagos = con aseguradora externa.")

    # --- Cajitas
    cajitas, caj_after = [], 0
    for fila, r in read_sheet(wb, "Cajitas"):
        cid = fk("Cajitas", fila, _str(r["cliente_id"]))
        f = _iso(r["fecha_abono"])
        tipo = _str(r["tipo_cajita"])
        if tipo not in ENUMS["tipo_cajita"]:
            raise SchemaError(f"Cajitas fila {fila}: tipo {tipo!r}")
        monto = _int(r["monto_abono_cop"])
        if monto <= 0:
            raise SchemaError(f"Cajitas fila {fila}: monto no positivo")
        if f > corte:
            caj_after += 1
            excluded.append(["Cajitas", fila, cid, f, f"fecha posterior al corte {corte}", tipo])
            continue
        cajitas.append([f"K{fila:05d}", cid, f, tipo, monto])
    rep.add("Cajitas", "fechas posteriores al corte", "aviso" if caj_after else "ok",
            f"{caj_after} abonos después de {corte} excluidos" if caj_after else "ninguno", caj_after)
    with_k = {k[1] for k in cajitas} | {e[2] for e in excluded if e[0] == "Cajitas"}
    flag_mismatch = sum((c[8] == "Sí") != (c[0] in with_k) for c in clientes)
    rep.add("Cajitas", "consistencia con Clientes.cajitas_activas", "ok" if not flag_mismatch else "aviso",
            f"{flag_mismatch} clientes con la bandera distinta a la tabla", flag_mismatch)

    # --- Uso_App
    uso, seen = [], set()
    for fila, r in read_sheet(wb, "Uso_App"):
        cid = fk("Uso_App", fila, _str(r["cliente_id"]))
        if cid in seen:
            raise SchemaError(f"Uso_App fila {fila}: cliente repetido")
        seen.add(cid)
        vals = [_int(r[c]) for c in SCHEMAS["Uso_App"][1:7]]
        lat = _str(r["latencia_eventos"])
        if lat not in ENUMS["latencia_eventos"]:
            raise SchemaError(f"Uso_App fila {fila}: latencia {lat!r}")
        if vals[3] + vals[4] != vals[2]:
            raise SchemaError(f"Uso_App fila {fila}: abiertas + ignoradas ≠ enviadas")
        uso.append([cid, *vals, lat])
    missing = ids - seen
    rep.add("Uso_App", "un registro por cliente", "ok" if not missing else "error",
            f"{len(uso)} registros; faltan {len(missing)}", len(uso))
    caps = Counter(u[6] for u in uso)
    at_cap = sum(u[3] >= u[6] for u in uso)
    rep.add("Uso_App", "tope de contactos", "ok",
            f"max_contactos_mes: {dict(caps)}; {at_cap} clientes ya enviaron ≥ tope en 30 días", at_cap)

    # --- Seguros_Mensual
    sm = []
    for fila, r in read_sheet(wb, "Seguros_Mensual"):
        ramo = _str(r["ramo"])
        if ramo not in ENUMS["ramo"]:
            raise SchemaError(f"Seguros_Mensual fila {fila}: ramo {ramo!r}")
        sm.append([_iso(r["mes"])[:7], ramo, _int(r["polizas_emitidas"]), _int(r["renovaciones"]),
                   _int(r["prima_promedio_cop"]), _fmt_num(_num(r["comision_pct"])),
                   _fmt_num(_num(r["persistencia_6m"])), _fmt_num(_num(r["persistencia_12m"])),
                   _fmt_num(_num(r["cancelacion_pct"]))])
    meses = sorted({s[0] for s in sm})
    rep.add("Seguros_Mensual", "cobertura", "ok", f"{len(meses)} meses ({meses[0]} a {meses[-1]}) × "
            f"{len({s[1] for s in sm})} ramos", len(sm))
    nulls = Counter(s[1] for s in sm if s[7] == "")
    if nulls:
        rep.add("Seguros_Mensual", "persistencia vacía", "aviso",
                "sin persistencia_12m: " + ", ".join(f"{k} ({v} meses)" for k, v in sorted(nulls.items())))

    # --- Elegibilidad
    eleg = [[_str(r[c]) for c in SCHEMAS["Elegibilidad"]] for _, r in read_sheet(wb, "Elegibilidad")]
    rep.add("Elegibilidad", "ramos", "ok" if {e[0] for e in eleg} == ENUMS["ramo"] else "error",
            ", ".join(f"{e[0]} ({e[1]}, {e[2]})" for e in eleg), len(eleg))
    rep.add("Elegibilidad", "rangos de edad", "aviso",
            "Clientes trae la edad en rangos; «56+» no se puede comparar con los topes de 60 (Desempleo) y 65 (Vida)")

    # --- Campañas previas: se conserva la cifra reportada y se recalculan las tasas
    camp, mismatches = [], []
    for fila, r in read_sheet(wb, "Campanas_Previas"):
        el, im, it, ve = (_int(r[c]) for c in ("elegibles", "impactados", "interesados", "ventas"))
        rep_conv = _num(r["conversion_impactados"])
        vi, vel, ii = ve / im, ve / el, it / im
        nota = ""
        if abs(rep_conv - vi) > 5e-4:
            if abs(rep_conv - vel) <= 5e-4:
                nota = "conversion_impactados reportada = ventas/elegibles, no ventas/impactados"
            else:
                nota = "conversion_impactados reportada no coincide con ningún cociente"
            mismatches.append(_str(r["campana"]))
        camp.append([_str(r["campana"]), _str(r["canal"]), el, im, it, ve, _fmt_num(rep_conv),
                     _fmt_num(round(vi, 4)), _fmt_num(round(vel, 4)), _fmt_num(round(ii, 4)), nota])
    rep.add("Campanas_Previas", "conversion_impactados", "aviso" if mismatches else "ok",
            (f"{len(mismatches)} de {len(camp)} filas: la columna es ventas/elegibles. Se conserva como "
             f"«conversion_reportada» y se agregan los cocientes recalculados.") if mismatches else "coincide",
            len(mismatches))

    # --- Diccionario: nombres de columnas documentados que no existen
    dic = [(_str(r["pestana"]), _str(r["campo"]))
           for _, r in read_sheet(wb, "Diccionario", ["pestana", "campo", "descripcion", "tipo", "ejemplo"])]
    ghost = [f"{p}.{c}" for p, c in dic if p in SCHEMAS and c not in SCHEMAS[p]]
    rep.add("Diccionario", "campos documentados", "aviso" if ghost else "ok",
            ("no existen en la pestaña: " + ", ".join(ghost)) if ghost else "todos existen", len(ghost))

    # --- escribir
    out.mkdir(parents=True, exist_ok=True)
    files = {}
    files["clientes.csv"] = write_csv(out / "clientes.csv", SCHEMAS["Clientes"], sorted(clientes))
    files["movimientos.csv"] = write_csv(out / "movimientos.csv", ["mov_id", *SCHEMAS["Movimientos"]],
                                         sorted(movs, key=lambda x: (x[1], x[2], x[0])))
    files["pagos_seguros.csv"] = write_csv(out / "pagos_seguros.csv", ["pago_id", *SCHEMAS["Pagos_Seguros"]],
                                           sorted(pagos, key=lambda x: (x[1], x[2], x[0])))
    files["cajitas.csv"] = write_csv(out / "cajitas.csv", ["abono_id", *SCHEMAS["Cajitas"]],
                                     sorted(cajitas, key=lambda x: (x[1], x[2], x[0])))
    files["uso_app.csv"] = write_csv(out / "uso_app.csv", SCHEMAS["Uso_App"], sorted(uso))
    files["seguros_mensual.csv"] = write_csv(out / "seguros_mensual.csv", SCHEMAS["Seguros_Mensual"], sm)
    files["elegibilidad.csv"] = write_csv(out / "elegibilidad.csv", SCHEMAS["Elegibilidad"], eleg)
    files["campanas_previas.csv"] = write_csv(
        out / "campanas_previas.csv",
        ["campana", "canal", "elegibles", "impactados", "interesados", "ventas", "conversion_reportada",
         "conv_ventas_impactados", "conv_ventas_elegibles", "conv_interesados_impactados", "qa_nota"], camp)
    files["juridico.csv"] = write_csv(out / "juridico.csv", SCHEMAS["Juridico_Entrega"],
                                      [[_str(r["tema"]), _str(r["supuesto_simulado"]), _str(r["nota"])]
                                       for r in jur.values()])
    files["excluidos.csv"] = write_csv(out / "excluidos.csv",
                                       ["pestana", "fila_excel", "cliente_id", "fecha", "razon", "detalle"],
                                       sorted(excluded, key=lambda x: (x[0], x[1])))

    manifest = {
        "fuente": {"archivo": src.name, "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw)},
        "ingesta_version": INGEST_VERSION,
        "periodo": {"inicio": inicio, "corte": corte},
        "uso": str(leeme.get("Uso", "")),
        "restriccion": str(leeme.get("Restricción", "")),
        "filas": {"clientes": len(clientes), "movimientos": len(movs), "pagos_seguros": len(pagos),
                  "cajitas": len(cajitas), "uso_app": len(uso), "seguros_mensual": len(sm),
                  "elegibilidad": len(eleg), "campanas_previas": len(camp), "excluidos": len(excluded)},
        "chequeos": rep.checks,
        "archivos": dict(sorted(files.items())),
    }
    if rep.errors:
        raise SchemaError("Chequeos con error:\n" + "\n".join(f"  - {e['pestana']}: {e['detalle']}" for e in rep.errors))
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out / "CALIDAD.md").write_text(quality_md(manifest), encoding="utf-8")
    return manifest


def quality_md(m: dict) -> str:
    icon = {"ok": "OK", "aviso": "AVISO", "error": "ERROR"}
    lines = [
        "# Calidad de la base del reto",
        "",
        "Archivo generado por `scripts/ingest_reto.py`. No editar a mano.",
        "",
        f"- Fuente: `{m['fuente']['archivo']}` · sha256 `{m['fuente']['sha256'][:16]}…`",
        f"- Periodo: {m['periodo']['inicio']} a {m['periodo']['corte']} (fecha de corte)",
        f"- Uso: {m['uso']}. {m['restriccion']}",
        "",
        "| Tabla | Filas |",
        "|---|---:|",
        *[f"| {k} | {v:,} |".replace(",", ".") for k, v in m["filas"].items()],
        "",
        "| Pestaña | Chequeo | Estado | Detalle |",
        "|---|---|---|---|",
        *[f"| {c['pestana']} | {c['chequeo']} | {icon[c['estado']]} | {c['detalle']} |" for c in m["chequeos"]],
        "",
    ]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--src", type=Path, default=DEFAULT_SRC)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args(argv)
    try:
        m = ingest(args.src, args.out)
    except SchemaError as e:
        print(f"INGESTA DETENIDA: {e}", file=sys.stderr)
        return 1
    print(f"Fuente  {m['fuente']['archivo']}  sha256={m['fuente']['sha256'][:16]}…")
    print(f"Periodo {m['periodo']['inicio']} → corte {m['periodo']['corte']}")
    for k, v in m["filas"].items():
        print(f"  {k:<18}{v:>8,}".replace(",", "."))
    for c in m["chequeos"]:
        print(f"  [{c['estado'].upper():5}] {c['pestana']} · {c['chequeo']}: {c['detalle']}")
    print(f"Salida  {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
